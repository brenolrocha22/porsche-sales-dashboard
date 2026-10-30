"""
Prepara a base do dashboard a partir da saída do agente de tratamento.

1. Fica só com o ID e as colunas *Sanitized* (nome de cliente e de vendedor não vão
   para o HTML, porque a base fica embutida na página publicada).
2. Converte preço, ano e quilometragem em número; data INVALID vira vazio.
3. Deriva a família do modelo (911, 718, Cayenne, Macan, Panamera, Taycan).
4. Salva como valores em .xlsx e como .json (o JSON é o que vai embutido no index.html).

Uso:
    python prepare_base.py data/porsche_sanitized.xlsx data/porsche_dashboard_base.xlsx data/porsche_dashboard_base.json
"""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

INVALID = "INVALID"

# coluna de saída do agente -> nome curto usado no dashboard
COLUNAS = {
    "sale_id": "id",
    "SaleDateSanitized": "data",
    "PorscheModelSanitized": "modelo",
    "ModelYearSanitized": "ano",
    "SalesPriceSanitized": "preco",
    "VehicleMileageSanitized": "milhas",
    "PayMethodSanitized": "pagamento",
    "CitySanitized": "cidade",
    "StateSanitized": "estado",
    "DeliveryStatusSanitized": "entrega",
}
NUMERICAS = {"ano": int, "preco": float, "milhas": int}


def familia(modelo: str) -> str:
    for prefixo in ("911", "718", "Cayenne", "Macan", "Panamera", "Taycan"):
        if modelo.startswith(prefixo):
            return prefixo
    return "Outros"


def preparar(entrada: Path) -> list[dict]:
    ws = load_workbook(entrada, data_only=True).active
    linhas = list(ws.iter_rows(values_only=True))
    idx = {h: i for i, h in enumerate(linhas[0])}
    registros = []
    for r in linhas[1:]:
        reg = {}
        for origem, destino in COLUNAS.items():
            v = r[idx[origem]]
            if v == INVALID:
                v = None
            elif destino in NUMERICAS and v is not None:
                v = NUMERICAS[destino](float(v))
            reg[destino] = v
        reg["familia"] = familia(reg["modelo"] or "")
        registros.append(reg)
    return registros


def salvar_xlsx(registros: list[dict], saida: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "base_dashboard"
    cab = list(registros[0].keys())
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
    for reg in registros:
        ws.append([reg[k] for k in cab])
    ws.freeze_panes = "A2"
    wb.save(saida)


def main() -> int:
    entrada, saida_xlsx, saida_json = (Path(a) for a in sys.argv[1:4])
    registros = preparar(entrada)
    salvar_xlsx(registros, saida_xlsx)
    saida_json.write_text(json.dumps(registros, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    print(f"{len(registros)} registros -> {saida_xlsx.name}, {saida_json.name}")
    for campo in ("familia", "modelo", "ano", "pagamento", "estado", "cidade", "entrega"):
        print(f"\n{campo}: {Counter(r[campo] for r in registros).most_common()}")
    precos = sorted(r["preco"] for r in registros if r["preco"] is not None)
    milhas = sorted(r["milhas"] for r in registros if r["milhas"] is not None)
    print(f"\npreço: min {precos[0]:,.2f} | mediana {precos[len(precos)//2]:,.2f} | max {precos[-1]:,.2f} | total {sum(precos):,.2f}")
    print(f"milhas: min {milhas[0]} | max {milhas[-1]}")
    print(f"datas válidas: {sum(1 for r in registros if r['data'])} | faixa {min(r['data'] for r in registros if r['data'])} a {max(r['data'] for r in registros if r['data'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
